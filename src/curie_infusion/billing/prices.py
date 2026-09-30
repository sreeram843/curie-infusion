"""CMS price files, read from a local directory (default `data/cms`, gitignored).

- ASP: the quarterly Medicare Part B drug payment limit file (public). We read its "section 508"
  CSV: HCPCS code, short description, dosage per billing unit, payment limit per billing unit.
- OPPS Addendum B: hospital outpatient payment rates. It is distributed behind the AMA CPT license,
  so the user downloads it. We read only the code, status indicator and payment rate; CPT
  descriptors are never stored or shown.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

# CMS dosage unit -> (dimension, factor to the dimension's base unit)
DOSE_UNITS = {
    "mg": ("mass", 1.0), "mcg": ("mass", 0.001), "gm": ("mass", 1000.0), "g": ("mass", 1000.0),
    "grams": ("mass", 1000.0), "gram": ("mass", 1000.0),
    "unit": ("unit", 1.0), "units": ("unit", 1.0),
    "meq": ("meq", 1.0), "meq.": ("meq", 1.0),
    "mmol": ("mmol", 1.0),
    "ml": ("volume", 1.0), "cc": ("volume", 1.0),
}
ADDENDUM_B = re.compile(r"addendum[\s_-]*b", re.I)
DOSAGE = re.compile(r"^\s*(\d*\.?\d+)\s*([a-z.]+)\s*$", re.I)


def parse_dosage(text: str) -> tuple[float, str] | None:
    """'2 MEQ' -> (2.0, 'meq'); '0.1 MG' -> (0.1, 'mg'). Compound dosages ('10 MG/3MG') -> None."""
    m = DOSAGE.match(text or "")
    if not m or m.group(2).lower() not in DOSE_UNITS:
        return None
    return float(m.group(1)), m.group(2).lower()


def billing_units(amount: float, unit: str, dosage: str) -> float | None:
    """Exact (unrounded) number of HCPCS billing units in `amount` `unit`, or None if incompatible."""
    parsed = parse_dosage(dosage)
    have = DOSE_UNITS.get((unit or "").lower())
    if not parsed or not have:
        return None
    size, size_unit = parsed
    want = DOSE_UNITS[size_unit]
    if have[0] != want[0]:
        return None
    return amount * have[1] / (size * want[1])


@dataclass
class PriceFile:
    source: str
    effective: str
    codes: dict[str, dict] = field(default_factory=dict)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):  # CMS ships Windows-1252
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def load_asp(root: Path) -> PriceFile | None:
    files = sorted(
        p for p in Path(root).rglob("*.csv")
        if "payment limit file" in p.name.lower() and "not payable" not in p.name.lower()
    )
    if not files:
        return None
    path = files[-1]
    rows = list(csv.reader(_read_text(path).splitlines()))
    effective = next((r[0] for r in rows if r and r[0].lower().startswith("effective")), "")
    header = next(i for i, r in enumerate(rows) if r and r[0].strip() == "HCPCS Code")
    codes = {}
    for r in rows[header + 1:]:
        if len(r) < 4 or not r[0].strip():
            continue
        try:
            limit = float(r[3])
        except ValueError:
            continue
        codes[r[0].strip()] = {"description": r[1].strip(), "dosage": r[2].strip(), "limit": limit}
    return PriceFile(path.name, effective, codes)


def _tables_in(path: Path) -> list[tuple[str, list[list]]]:
    """(name, rows) for every CSV/XLSX in a file or zip."""
    if path.suffix.lower() == ".zip":
        out = []
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                data = z.read(name)
                if name.lower().endswith(".csv"):
                    out.append((name, list(csv.reader(io.StringIO(data.decode("cp1252", "replace"))))))
                elif name.lower().endswith(".xlsx"):
                    out.append((name, _xlsx_rows(io.BytesIO(data))))
        return out
    if path.suffix.lower() == ".csv":
        return [(path.name, list(csv.reader(_read_text(path).splitlines())))]
    if path.suffix.lower() == ".xlsx":
        return [(path.name, _xlsx_rows(path))]
    return []


def _xlsx_rows(src) -> list[list]:
    from openpyxl import load_workbook

    ws = load_workbook(src, read_only=True, data_only=True).worksheets[0]
    return [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]


def load_opps(root: Path, wanted: set[str]) -> PriceFile | None:
    """Payment rate and status indicator for the `wanted` codes from an Addendum B file, if present."""
    candidates = sorted(
        p for p in Path(root).rglob("*")
        if p.is_file() and ADDENDUM_B.search(p.name) and p.suffix.lower() in {".zip", ".csv", ".xlsx"}
    )
    for path in reversed(candidates):
        for name, rows in _tables_in(path):
            for i, r in enumerate(rows[:30]):
                cells = [c.strip().lower() for c in r]
                code_col = next((j for j, c in enumerate(cells) if c.startswith("hcpcs")), None)
                rate_col = next((j for j, c in enumerate(cells) if "payment rate" in c), None)
                si_col = next((j for j, c in enumerate(cells) if c in {"si", "status indicator"}), None)
                if code_col is None or rate_col is None:
                    continue
                codes = {}
                for row in rows[i + 1:]:
                    if len(row) <= max(code_col, rate_col) or row[code_col].strip() not in wanted:
                        continue
                    rate = re.sub(r"[$,\s]", "", row[rate_col])
                    codes[row[code_col].strip()] = {
                        "rate": float(rate) if rate else None,
                        "si": row[si_col].strip() if si_col is not None and len(row) > si_col else "",
                    }
                return PriceFile(f"{path.name}:{name}", "", codes)
    return None
